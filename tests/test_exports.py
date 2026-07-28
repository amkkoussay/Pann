"""
Unit tests for the PDF + Excel export service.
"""
import io
import zipfile
import pytest

from openpyxl import load_workbook

from models import db, Customer, Product, ProductionOrder
from services import ProfitabilityService
from exports import (
    customer_statement_pdf, production_order_invoice_pdf,
    rows_to_xlsx, profitability_xlsx, reorder_suggestions_xlsx,
)


def test_customer_statement_pdf_is_valid_pdf(db_session):
    cust = db_session.query(Customer).first()
    buf = customer_statement_pdf(cust.id)
    assert isinstance(buf, io.BytesIO)
    head = buf.read(5)
    buf.seek(0)
    assert head == b'%PDF-'


def test_customer_statement_pdf_404_on_missing(db_session):
    with pytest.raises(ValueError):
        customer_statement_pdf(99999)


def test_production_invoice_pdf_is_valid_pdf(db_session):
    bread = db_session.query(Product).filter_by(name='Sourdough').one()
    order = ProductionOrder(product_id=bread.id, planned_qty=10,
                            actual_qty=10, status='completed')
    db_session.add(order); db_session.commit()
    buf = production_order_invoice_pdf(order.id)
    head = buf.read(5)
    buf.seek(0)
    assert head == b'%PDF-'


def test_xlsx_round_trip():
    rows = [
        {'name': 'Flour', 'qty': 100, 'cost': 2.5},
        {'name': 'Sugar', 'qty': 50, 'cost': 2.0},
    ]
    buf = rows_to_xlsx(rows, sheet_name='Inventory')
    # XLSX is a zip → check PK header bytes.
    assert buf.read(2) == b'PK'
    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active
    assert ws.title == 'Inventory'
    headers = [c.value for c in ws[1]]
    assert headers == ['name', 'qty', 'cost']
    assert ws.cell(row=2, column=1).value == 'Flour'
    assert ws.cell(row=3, column=2).value == 50


def test_profitability_xlsx_export(db_session):
    svc = ProfitabilityService()
    buf = profitability_xlsx(svc.analyze_all())
    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active
    assert ws.title == 'Profitability'
    headers = [c.value for c in ws[1]]
    assert 'margin' in headers
    assert 'net_profit' in headers


def test_reorder_xlsx_export(app):
    with app.app_context():
        from services import SmartInventoryService
        svc = SmartInventoryService()
        buf = reorder_suggestions_xlsx(svc.analyze_all())
        buf.seek(0)
        wb = load_workbook(buf)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert 'recommended_order_qty' in headers
        assert 'urgency' in headers


def test_xlsx_handles_empty_rows():
    buf = rows_to_xlsx([], sheet_name='Empty')
    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == '(no data)'
